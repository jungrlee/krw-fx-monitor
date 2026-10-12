#!/usr/bin/env python3
"""
KRW FX Monitor - builds a formula-driven Excel workbook for a USD/KRW desk.

Usage
    python krw_fx_monitor.py                      # try live sources, fall back to the bundled snapshot
    python krw_fx_monitor.py --offline            # bundled snapshot only (no internet needed)
    python krw_fx_monitor.py --ecos-key YOUR_KEY  # also pull Bank of Korea ECOS series (optional)
    python krw_fx_monitor.py --out My_Report.xlsx --start 2000-01-01

Data sources (all free; see the Methodology sheet of the output for what was actually used)
    FRED (St. Louis Fed, no key):  DEXKOUS, DEXJPUS, DEXCHUS (Fed H.10 daily), DTWEXBGS (broad dollar index),
                                   VIXCLS, DGS10, SOFR, IR3TIB01KRM156N / IR3TIB01USM156N (OECD 3-month rates)
    GitHub mirror of Fed H.10 (datasets/exchange-rates), used if FRED is unreachable
    World Bank API (no key):       GDP growth, CPI inflation, Korea current account (% of GDP)
    Bank of Korea ECOS (optional free key from ecos.bok.or.kr): USD/KRW base rate, CD 91-day, base rate

    KRX foreign net buying via pykrx (optional: pip install pykrx), or your own CSV with --krx-csv
    Flow data from ECOS/KRX are cached in data/ after a successful download, so --offline reuses them.

Requires: Python 3.9+, pandas, openpyxl  (pip install pandas openpyxl; optional: pykrx)
"""
import argparse
import io
import json
import os
import sys
import urllib.request
from datetime import date, datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SNAP = os.path.join(HERE, "data")
UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0 Safari/537.36"}

# ----------------------------------------------------------------------------------------------
# Data layer
# ----------------------------------------------------------------------------------------------
LOG = []  # (series, description, source, obs, first, last, note)


def _get(url, timeout=30):
    import ssl
    try:  # macOS python.org builds often lack root certificates; certifi supplies them
        import certifi
        ctx = ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        ctx = None
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout, context=ctx) as r:
        return r.read()


FRED_DOWN = []  # set after the first timeout so the run doesn't wait on every series


def fred(series_id, start):
    """FRED series. Uses the official API if FRED_API_KEY is set (free key), else the public CSV endpoint."""
    key = os.getenv("FRED_API_KEY")
    if key:
        js = json.loads(_get(f"https://api.stlouisfed.org/fred/series/observations?series_id={series_id}"
                             f"&api_key={key}&file_type=json&observation_start={start}", 60))
        s = pd.Series({pd.Timestamp(o["date"]): pd.to_numeric(o["value"], errors="coerce") for o in js["observations"]})
        return s.dropna().sort_index()
    if FRED_DOWN:
        raise TimeoutError("skipped: FRED timed out earlier in this run (set FRED_API_KEY)")
    try:
        raw = _get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series_id}&cosd={start}", 45)
    except Exception:
        FRED_DOWN.append(1)
        raise
    df = pd.read_csv(io.BytesIO(raw))
    df.columns = ["date", "value"]
    df["date"] = pd.to_datetime(df["date"])
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    return df.dropna().set_index("date")["value"]


def h10_mirror(country, start):
    """Fed H.10 daily rates mirrored on GitHub (datasets/exchange-rates). Matches FRED exactly in tests."""
    raw = _get("https://raw.githubusercontent.com/datasets/exchange-rates/main/data/daily.csv", 60)
    df = pd.read_csv(io.BytesIO(raw))
    df = df[df["Country"] == country].copy()
    df["Date"] = pd.to_datetime(df["Date"])
    s = pd.to_numeric(df.set_index("Date")["Exchange rate"], errors="coerce").dropna()
    return s[s.index >= start]


def snapshot(name, start=None, col="value"):
    df = pd.read_csv(os.path.join(SNAP, name), parse_dates=[0])
    s = df.set_index(df.columns[0])[col] if col in df.columns else df.set_index(df.columns[0]).iloc[:, 0]
    return s[s.index >= start] if start else s


def worldbank(indicator, country):
    raw = _get(f"https://api.worldbank.org/v2/country/{country}/indicator/{indicator}?format=json&per_page=200", 40)
    js = json.loads(raw)
    rows = {int(r["date"]): r["value"] for r in js[1] if r["value"] is not None}
    return pd.Series(rows).sort_index()


def ecos(key, stat, freq, item, start, end):
    url = (f"https://ecos.bok.or.kr/api/StatisticSearch/{key}/json/kr/1/100000/"
           f"{stat}/{freq}/{start}/{end}/{item}")
    js = json.loads(_get(url, 40))
    rows = js["StatisticSearch"]["row"]
    fmt = "%Y%m%d" if freq == "D" else "%Y%m"
    s = pd.Series({datetime.strptime(r["TIME"], fmt): float(r["DATA_VALUE"]) for r in rows})
    return s.sort_index()


# ---------------------------------------------------------------- flows: balance of payments (ECOS)
# Item codes differ between ECOS table versions, so items are found by name in the ECOS item list.
# Override with flows_config.json next to this script: {"ca": ["301Y013", "000000"], ...}
BOP_TABLES = ["301Y013", "301Y017", "301Y014"]
BOP_ITEMS = {  # key: (description, regexes tried against "parent / child" item names, English)
    "ca":   ("Current account",                         [r"^current account$"]),
    "goods": ("Goods balance",                          [r"^(current account / )?goods( balance)?$"]),
    "di_a": ("Direct investment, assets (outward)",     [r"direct investment.*assets", r"direct investment abroad"]),
    "di_l": ("Direct investment, liabilities (inward)", [r"direct investment.*liabilities", r"direct investment in korea"]),
    "pi_a": ("Portfolio investment, assets (outward)",  [r"portfolio investment.*assets"]),
    "pi_l": ("Portfolio investment, liabilities (inward)", [r"portfolio investment.*liabilities"]),
}


def ecos_items(key, stat):
    js = json.loads(_get(f"https://ecos.bok.or.kr/api/StatisticItemList/{key}/json/en/1/2000/{stat}", 40))
    rows = js["StatisticItemList"]["row"]
    names = {r["ITEM_CODE"]: r["ITEM_NAME"].strip() for r in rows}
    out = []
    for r in rows:
        path, p, seen = [r["ITEM_NAME"].strip()], r.get("P_ITEM_CODE"), set()
        while p and p in names and p not in seen:
            seen.add(p)
            path.insert(0, names[p])
            p = next((x.get("P_ITEM_CODE") for x in rows if x["ITEM_CODE"] == p), None)
        out.append((r["ITEM_CODE"], " / ".join(path), r.get("CYCLE", "")))
    return out


def find_bop_codes(key):
    import re
    cfg_path = os.path.join(HERE, "flows_config.json")
    codes = json.load(open(cfg_path)) if os.path.exists(cfg_path) else {}
    for stat in BOP_TABLES:
        missing = [k for k in BOP_ITEMS if k not in codes]
        if not missing:
            break
        try:
            items = [it for it in ecos_items(key, stat) if "M" in it[2] or not it[2]]
        except Exception:  # noqa: BLE001
            continue
        for k in missing:
            for pat in BOP_ITEMS[k][1]:
                hits = [it for it in items if re.search(pat, it[1].lower())]
                if hits:
                    best = min(hits, key=lambda it: len(it[1]))
                    codes[k] = [stat, best[0], best[1]]
                    break
    return codes


def load_bop(args):
    """Monthly balance of payments, US$ mn, from ECOS. Falls back to the last successful download (cache)."""
    cache = os.path.join(SNAP, "cache_bop_monthly.csv")
    if args.ecos_key and not args.offline:
        try:
            codes = find_bop_codes(args.ecos_key)
            end = date.today().strftime("%Y%m")
            cols = {}
            for k, v in codes.items():
                stat, item = v[0], v[1]
                try:
                    cols[k] = ecos(args.ecos_key, stat, "M", item, "199501", end)
                    label = v[2] if len(v) > 2 else item
                    print(f"  BoP {k:<6} ECOS {stat}/{item:<10} '{label}'  {len(cols[k])} obs")
                    LOG.append((f"BoP {k}", BOP_ITEMS[k][0] + f", US$ mn (matched '{label}')", f"ECOS {stat}/{item}",
                                len(cols[k]), cols[k].index.min(), cols[k].index.max(), ""))
                except Exception as e:  # noqa: BLE001
                    print(f"  BoP {k:<6} ECOS {stat}/{item} failed ({type(e).__name__})")
            for k in BOP_ITEMS:
                if k not in codes:
                    print(f"  BoP {k:<6} no matching ECOS item found (see --ecos-list)")
                    LOG.append((f"BoP {k}", BOP_ITEMS[k][0], "unavailable", 0, None, None, "no ECOS item matched"))
            if cols:
                df = pd.DataFrame(cols)
                df.index.name = "month"
                df.to_csv(cache)
                return df
        except Exception as e:  # noqa: BLE001
            print(f"  BoP    ECOS failed ({type(e).__name__})")
    if os.path.exists(cache):
        df = pd.read_csv(cache, parse_dates=["month"]).set_index("month")
        LOG.append(("BoP", "Balance of payments, US$ mn", "cache of last ECOS download", len(df), df.index.min(), df.index.max(), ""))
        print(f"  BoP    cache of last ECOS download            {len(df):>6} obs  to {df.index.max().date()}")
        return df
    LOG.append(("BoP", "Balance of payments, US$ mn", "unavailable", 0, None, None, "needs --ecos-key (free at ecos.bok.or.kr)"))
    print("  BoP    unavailable (needs --ecos-key)")
    return None


def krx_login_check():
    """Log in to KRX once and stop with KRX's own message if it fails (avoids repeated failed logins)."""
    import requests
    from pykrx.website.comm import auth
    s = requests.Session()
    auth.warmup_krx_session(s)
    payload = {"mbrNm": "", "telNo": "", "di": "", "certType": "",
               "mbrId": os.getenv("KRX_ID", ""), "pw": os.getenv("KRX_PW", "")}
    headers = {"User-Agent": auth.USER_AGENT, "Referer": auth.LOGIN_PAGE}
    js = s.post(auth.LOGIN_URL, data=payload, headers=headers, timeout=15).json()
    if js.get("_error_code") == "CD011":  # already logged in elsewhere
        payload["skipDup"] = "Y"
        js = s.post(auth.LOGIN_URL, data=payload, headers=headers, timeout=15).json()
    code, msg = js.get("_error_code", ""), js.get("_error_message", "")
    if code != "CD001":
        print(f"  Foreign  KRX login refused: {code} {msg}")
        raise RuntimeError(f"KRX login refused ({code})")
    print("  Foreign  KRX login OK")


def krx_foreign(start):
    """Daily foreign net buying (KRW bn) on KOSPI and KOSDAQ via pykrx (pip install pykrx)."""
    krx_login_check()
    from pykrx import stock  # imported lazily: optional dependency
    out = {}
    start, end = pd.Timestamp(start).date(), date.today()
    for mkt in ("KOSPI", "KOSDAQ"):
        parts = []
        for y in range(start.year, end.year + 1):
            f, t = max(start, date(y, 1, 1)), min(end, date(y, 12, 31))
            df = stock.get_market_trading_value_by_date(f.strftime("%Y%m%d"), t.strftime("%Y%m%d"), mkt)
            if len(df):
                parts.append(df["외국인합계"].astype(float) / 1e9)
        out[mkt.lower()] = pd.concat(parts)
    df = pd.DataFrame(out)
    df.index = pd.to_datetime(df.index)
    return df.sort_index()


def load_krx(args):
    cache = os.path.join(SNAP, "cache_krx_foreign_daily.csv")
    tries = []
    if args.krx_csv:
        tries.append(("your CSV " + os.path.basename(args.krx_csv),
                      lambda: pd.read_csv(args.krx_csv, parse_dates=[0], index_col=0)))
    if not args.offline:
        if os.getenv("KRX_ID") and os.getenv("KRX_PW"):
            tries.append(("KRX via pykrx", lambda: krx_foreign(args.flows_start)))
        else:
            print("  Foreign  KRX skipped: KRX now requires a login. Create a free account at data.krx.co.kr, then\n"
                  "           export KRX_ID=... and export KRX_PW=... (pykrx >= 1.2.9 uses them).")
    if os.path.exists(cache):
        tries.append(("cache of last KRX download", lambda: pd.read_csv(cache, parse_dates=[0], index_col=0)))
    for label, fn in tries:
        try:
            df = fn()
            df.columns = [c.lower() for c in df.columns]
            df = df[[c for c in ("kospi", "kosdaq") if c in df.columns]].dropna(how="all")
            if len(df):
                if "cache" not in label:
                    df.index.name = "date"
                    df.to_csv(cache)
                LOG.append(("Foreign equity", "Foreign investors' net buying, KRW bn, daily (KRX)", label, len(df),
                            df.index.min(), df.index.max(), ""))
                print(f"  Foreign  {label:<40} {len(df):>6} obs  to {df.index.max().date()}")
                return df
        except Exception as e:  # noqa: BLE001
            print(f"  Foreign  {label:<40} failed ({type(e).__name__}: {str(e)[:60]})")
    LOG.append(("Foreign equity", "Foreign investors' net buying, KRW bn, daily (KRX)", "unavailable", 0, None, None,
                "KRX requires login: set KRX_ID/KRX_PW (pip install -U pykrx), or pass --krx-csv"))
    print("  Foreign  unavailable")
    return None


def load(name, desc, tries, offline):
    """tries: list of (source_label, callable). Snapshot entries are marked with 'snapshot' in the label."""
    for label, fn in tries:
        if offline and "snapshot" not in label:
            continue
        try:
            s = fn()
            if s is not None and len(s) > 0:
                LOG.append((name, desc, label, len(s), s.index.min(), s.index.max(), ""))
                last = s.index.max()
                last = last if isinstance(last, (int, np.integer)) else pd.Timestamp(last).date()
                print(f"  {name:<10} {label:<40} {len(s):>6} obs  to {last}")
                return s
        except Exception as e:  # noqa: BLE001 - any source failure falls through to the next source
            print(f"  {name:<10} {label:<40} failed ({type(e).__name__}: {str(e)[:70]})")
    LOG.append((name, desc, "unavailable", 0, None, None, "all sources failed"))
    print(f"  {name:<10} unavailable")
    return None


def collect(args):
    st = args.start
    print("Loading data ...")
    d = {}
    d["usdkrw"] = load("USD/KRW", "KRW per USD, Fed H.10 noon NY buying rate, daily", [
        ("FRED DEXKOUS", lambda: fred("DEXKOUS", st)),
        ("GitHub mirror of Fed H.10", lambda: h10_mirror("South Korea", st)),
        ("bundled snapshot (Fed H.10)", lambda: snapshot("usdkrw_daily.csv", st))], args.offline)
    d["usdjpy"] = load("USD/JPY", "JPY per USD, Fed H.10, daily", [
        ("FRED DEXJPUS", lambda: fred("DEXJPUS", st)),
        ("GitHub mirror of Fed H.10", lambda: h10_mirror("Japan", st)),
        ("bundled snapshot (Fed H.10)", lambda: snapshot("usdjpy_daily.csv", st))], args.offline)
    d["usdcny"] = load("USD/CNY", "CNY per USD, Fed H.10, daily", [
        ("FRED DEXCHUS", lambda: fred("DEXCHUS", st)),
        ("GitHub mirror of Fed H.10", lambda: h10_mirror("China", st)),
        ("bundled snapshot (Fed H.10)", lambda: snapshot("usdcny_daily.csv", st))], args.offline)
    d["dxy"] = load("USD broad", "Nominal broad US dollar index (Fed), daily", [
        ("FRED DTWEXBGS", lambda: fred("DTWEXBGS", st))], args.offline)
    d["vix"] = load("VIX", "CBOE VIX, daily close", [("FRED VIXCLS", lambda: fred("VIXCLS", st))], args.offline)
    d["ust10"] = load("UST 10Y", "US 10-year Treasury yield, %, daily", [("FRED DGS10", lambda: fred("DGS10", st))], args.offline)
    d["sofr"] = load("SOFR", "Secured overnight financing rate, %, daily", [("FRED SOFR", lambda: fred("SOFR", st))], args.offline)
    d["kr3m"] = load("KRW 3M", "Korea 3-month interbank rate, %, monthly (OECD)", [
        ("FRED IR3TIB01KRM156N", lambda: fred("IR3TIB01KRM156N", "1995-01-01")),
        ("bundled snapshot (OECD via FRED)", lambda: snapshot("rates_monthly.csv", col="krw_3m"))], args.offline)
    d["us3m"] = load("USD 3M", "US 3-month interbank rate, %, monthly (OECD)", [
        ("FRED IR3TIB01USM156N", lambda: fred("IR3TIB01USM156N", "1995-01-01")),
        ("bundled snapshot (OECD via FRED)", lambda: snapshot("rates_monthly.csv", col="usd_3m"))], args.offline)

    macro = {}
    snapm = pd.read_csv(os.path.join(SNAP, "macro_annual.csv")).set_index("year")
    for key, ind, ctry, col in [("kr_gdp", "NY.GDP.MKTP.KD.ZG", "KOR", "kr_gdp"), ("us_gdp", "NY.GDP.MKTP.KD.ZG", "USA", "us_gdp"),
                                ("kr_cpi", "FP.CPI.TOTL.ZG", "KOR", "kr_cpi"), ("us_cpi", "FP.CPI.TOTL.ZG", "USA", "us_cpi"),
                                ("kr_ca", "BN.CAB.XOKA.GD.ZS", "KOR", "kr_ca_gdp")]:
        macro[key] = load(key, f"World Bank {ind} ({ctry}), annual, %", [
            (f"World Bank API {ind}", lambda i=ind, c=ctry: worldbank(i, c)),
            ("bundled snapshot (World Bank)", lambda c=col: snapm[c].dropna())], args.offline)
    d["macro"] = macro

    d["ecos"] = {}
    if args.ecos_key and not args.offline:
        end = date.today().strftime("%Y%m%d")
        for key, desc, stat, item in [("bok_usdkrw", "BOK USD/KRW base rate (매매기준율), daily", "731Y001", "0000001"),
                                      ("bok_cd91", "CD 91-day yield, %, daily", "817Y002", "010502000"),
                                      ("bok_base", "BOK base rate, %, daily", "722Y001", "0101000")]:
            d["ecos"][key] = load(key, desc + f" (ECOS {stat}/{item})",
                                  [(f"ECOS {stat}", lambda s=stat, i=item: ecos(args.ecos_key, s, "D", i, "20000101", end))], False)
    print("Loading flows ...")
    d["bop"] = load_bop(args)
    d["krx"] = load_krx(args)
    return d


# ----------------------------------------------------------------------------------------------
# Workbook layer (all analytics are Excel formulas; Python only writes raw data and layout)
# ----------------------------------------------------------------------------------------------
FONT = "Arial"
NAVY, RED, GREY, INK = "1F3A5F", "9B2226", "808080", "1A1A1A"
F_IN = Font(name=FONT, size=9, color="0000FF")           # hardcoded inputs / raw data
F_FX = Font(name=FONT, size=9, color="000000")           # formulas
F_LINK = Font(name=FONT, size=9, color="008000")         # links to other sheets
F_HDR = Font(name=FONT, size=9, bold=True, color=INK)
F_B = Font(name=FONT, size=9, bold=True, color=INK)
F_T = Font(name=FONT, size=13, bold=True, color=INK)
F_S = Font(name=FONT, size=8, color=GREY)
F_SEC = Font(name=FONT, size=9, bold=True, color=INK)
FILL_HDR = PatternFill(fill_type=None)
FILL_SEC = PatternFill(fill_type=None)
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="000000")
HAIR = Side(style="hair", color="A6A6A6")
MED = Side(style="medium", color="000000")
B_BOT = Border(bottom=THIN)
PCT, PCT2, NUM1, NUM2, NUM4 = '0.0%;-0.0%;-', '0.00%;-0.00%;-', '#,##0.0', '#,##0.00', '0.0000'
SGN = '+0.00%;-0.00%;0.00%'


def hdr(ws, row, labels, col=1, widths=None):
    """Column headers: bold, thin rule underneath, numbers right-aligned (sell-side table style)."""
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t if t != "" else None)
        c.font, c.border = F_HDR, Border(bottom=THIN)
        c.alignment = Alignment(horizontal="left" if i == 0 else "right", vertical="bottom", wrap_text=True)
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(col + i)].width = w


def section_at(ws, row, col, ncol, text):
    """Section title: bold, medium rule underneath, no fill."""
    for c in range(col, col + ncol):
        ws.cell(row=row, column=c).border = Border(bottom=MED)
    ws.cell(row=row, column=col, value=text).font = F_SEC


def section(ws, row, text, ncol=8):
    section_at(ws, row, 1, ncol, text)



# ---------------------------------------------------------------- charts
from openpyxl.chart import BarChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun

C_NAVY, C_RED, C_GREY, C_GREEN, C_BLACK = "1F3A5F", "9B2226", "A6A6A6", "5B8C5A", "1A1A1A"


def _cp(sz, b=False, color="404040"):
    return CharacterProperties(sz=sz, b=b, solidFill=color, latin=DFont(typeface=FONT))


def _rich(sz, b=False, color="404040"):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=_cp(sz, b, color)), endParaRPr=_cp(sz, b, color))])


def _title(text):
    cp = _cp(1000, True, "1A1A1A")
    t = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])])))
    t.overlay = False
    return t


def _axes(ch, y_fmt=None, x_fmt=None, skip=None):
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.txPr = _rich(800)
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill="808080", w=6350))
        ax.majorTickMark = "none"
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E3E3E3", w=6350)))
    ch.x_axis.majorGridlines = None
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.tickLblPos = "low"
    if y_fmt:
        ch.y_axis.number_format = y_fmt
    if x_fmt:
        ch.x_axis.number_format = x_fmt
    if skip:
        ch.x_axis.tickLblSkip = skip
        ch.x_axis.tickMarkSkip = skip


def new_chart(kind, title, w=16, h=7.5):
    ch = LineChart() if kind == "line" else BarChart()
    ch.title = _title(title)
    ch.width, ch.height = w, h
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.roundedCorners = False
    if kind == "bar":
        ch.type, ch.gapWidth = "col", 40
    ch.legend = None
    return ch


def add_line(ch, ws, col, r0, r1, name, colour, width=1.25):
    s = Series(Reference(ws, min_col=col, min_row=r0, max_row=r1), title=name)
    s.graphicalProperties.line.solidFill = colour
    s.graphicalProperties.line.width = int(width * 12700)
    s.smooth = False
    ch.series.append(s)
    return s


def add_bars(ch, ws, col, r0, r1, name, colour):
    s = Series(Reference(ws, min_col=col, min_row=r0, max_row=r1), title=name)
    s.graphicalProperties.solidFill = colour
    s.graphicalProperties.line.noFill = True
    ch.series.append(s)
    return s


def legend_bottom(ch):
    from openpyxl.chart.legend import Legend
    ch.legend = Legend()
    ch.legend.position = "b"
    ch.legend.txPr = _rich(800)


def combo(primary, secondary, y2_fmt=None):
    """Overlay `secondary` (line) on `primary` with its own right-hand axis."""
    secondary.y_axis.axId = 200
    secondary.y_axis.crosses = "max"
    secondary.y_axis.delete = False
    secondary.y_axis.majorGridlines = None
    secondary.y_axis.txPr = _rich(800)
    secondary.y_axis.majorTickMark = "none"
    secondary.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    if y2_fmt:
        secondary.y_axis.number_format = y2_fmt
    primary += secondary
    return primary


def put(ws, ref, value, font=F_FX, fmt=None, comment=None, fill=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if comment:
        c.comment = Comment(comment, "KRW FX Monitor")
    if fill:
        c.fill = fill
    return c


def build_daily(wb, d):
    ws = wb.create_sheet("Daily")
    k = d["usdkrw"]
    idx = k.index
    cols = [("USD/KRW", k), ("USD/JPY", d["usdjpy"]), ("USD/CNY", d["usdcny"]),
            ("USD broad index", d["dxy"]), ("VIX", d["vix"]), ("UST 10Y %", d["ust10"]), ("SOFR %", d["sofr"])]
    if d["ecos"].get("bok_usdkrw") is not None:
        cols.append(("BOK USD/KRW", d["ecos"]["bok_usdkrw"]))
    if d["ecos"].get("bok_cd91") is not None:
        cols.append(("CD 91D %", d["ecos"]["bok_cd91"]))
    if d.get("krx") is not None:
        cols.append(("Foreign net buy (KRW bn)", d["krx"].sum(axis=1, min_count=1)))
    frame = pd.DataFrame(index=idx)
    for name, s in cols:
        frame[name] = s.reindex(idx) if s is not None else None
    # Layout: A Date | B USD/KRW | C KRW log ret | D 20D vol | E USD/JPY | F JPY log ret | G USD/CNY | H CNY log ret | I.. extras
    labels = ["Date", "USD/KRW", "KRW log return", "KRW 20D vol (ann.)", "USD/JPY", "JPY log return", "USD/CNY", "CNY log return"]
    extras = [c for c, _ in cols[3:]]
    hdr(ws, 1, labels + extras, widths=[11, 10, 11, 11, 10, 11, 10, 11] + [12] * len(extras))
    ws.freeze_panes = "B2"
    n = len(frame)
    for i, (dt, row) in enumerate(frame.iterrows()):
        r = i + 2
        ws.cell(r, 1, dt.to_pydatetime()).number_format = "yyyy-mm-dd"
        ws.cell(r, 1).font = F_IN
        for col, name, fmt in [(2, "USD/KRW", NUM2), (5, "USD/JPY", NUM2), (7, "USD/CNY", NUM4)]:
            v = row[name]
            c = ws.cell(r, col, None if pd.isna(v) else float(v))
            c.font, c.number_format = F_IN, fmt
        for j, name in enumerate(extras):
            v = row[name]
            c = ws.cell(r, 9 + j, None if pd.isna(v) else float(v))
            c.font, c.number_format = F_IN, NUM2
        if r > 2:
            ws.cell(r, 3, f'=IF(OR(B{r}="",B{r-1}=""),"",LN(B{r}/B{r-1}))').number_format = '0.00%'
            ws.cell(r, 6, f'=IF(OR(E{r}="",E{r-1}=""),"",LN(E{r}/E{r-1}))').number_format = '0.00%'
            ws.cell(r, 8, f'=IF(OR(G{r}="",G{r-1}=""),"",LN(G{r}/G{r-1}))').number_format = '0.00%'
        if r >= 22:
            ws.cell(r, 4, f"=STDEV(C{r-19}:C{r})*SQRT(252)").number_format = '0.0%'
        for col in (3, 4, 6, 8):
            ws.cell(r, col).font = F_FX
    extra_col = {name: get_column_letter(9 + j) for j, name in enumerate(extras)}
    return ws, n, extra_col


def build_rates(wb, d):
    ws = wb.create_sheet("Rates_Monthly")
    m = pd.concat([d["kr3m"].rename("k"), d["us3m"].rename("u")], axis=1).dropna()
    hdr(ws, 1, ["Month", "KRW 3M %", "USD 3M %", "Gap (pts)", "USD/KRW avg", "Implied 3M pts (won)", "Hedge carry (ann.)"],
        widths=[11, 10, 10, 10, 12, 14, 14])
    ws.freeze_panes = "B2"
    for i, (dt, row) in enumerate(m.iterrows()):
        r = i + 2
        put(ws, f"A{r}", pd.Timestamp(dt).to_pydatetime(), F_IN, "yyyy-mm")
        put(ws, f"B{r}", float(row.k), F_IN, NUM2)
        put(ws, f"C{r}", float(row.u), F_IN, NUM2)
        put(ws, f"D{r}", f"=B{r}-C{r}", F_FX, NUM2)
        put(ws, f"E{r}", f'=IFERROR(AVERAGEIFS(Daily!$B:$B,Daily!$A:$A,">="&A{r},Daily!$A:$A,"<"&EDATE(A{r},1)),"")', F_LINK, NUM1)
        put(ws, f"F{r}", f'=IF(E{r}="","",E{r}*((1+B{r}/100*91/365)/(1+C{r}/100*91/360)-1))', F_FX, NUM2)
        put(ws, f"G{r}", f'=IF(E{r}="","",(F{r}/E{r})*365/91)', F_FX, PCT2)
    return ws, len(m)


def build_macro(wb, d):
    ws = wb.create_sheet("Macro_Annual")
    mc = d["macro"]
    years = sorted(set().union(*[set(s.index) for s in mc.values() if s is not None]))
    years = [y for y in years if 1996 <= y]
    hdr(ws, 1, ["Year", "KR GDP %", "US GDP %", "Growth gap", "KR CPI %", "US CPI %", "Inflation gap", "KR current acct % GDP",
                "USD/KRW Dec avg", "USD/KRW chg (log)"], widths=[8, 10, 10, 10, 10, 10, 11, 14, 13, 13])
    ws.freeze_panes = "B2"
    first = 2
    for i, y in enumerate(years):
        r = first + i
        put(ws, f"A{r}", str(y), F_IN)
        for col, key in [("B", "kr_gdp"), ("C", "us_gdp"), ("E", "kr_cpi"), ("F", "us_cpi"), ("H", "kr_ca")]:
            s = mc.get(key)
            v = s.get(y) if s is not None else None
            put(ws, f"{col}{r}", None if v is None or pd.isna(v) else float(v), F_IN, NUM2)
        put(ws, f"D{r}", f'=IF(OR(B{r}="",C{r}=""),"",B{r}-C{r})', F_FX, NUM2)
        put(ws, f"G{r}", f'=IF(OR(E{r}="",F{r}=""),"",E{r}-F{r})', F_FX, NUM2)
        put(ws, f"I{r}", f'=IFERROR(AVERAGEIFS(Daily!$B:$B,Daily!$A:$A,">="&DATE(VALUE(A{r}),12,1),'
                          f'Daily!$A:$A,"<"&DATE(VALUE(A{r})+1,1,1)),"")', F_LINK, NUM1)
        put(ws, f"J{r}", f'=IF(OR(I{r}="",I{r-1}="",ROW()={first}),"",LN(I{r}/I{r-1}))' if r > first else "", F_FX, PCT)
    last = first + len(years) - 1
    r = last + 2
    put(ws, f"A{r}", "R² vs. USD/KRW change", F_B)
    for col in ("D", "G", "H"):
        put(ws, f"{col}{r}", f"=IFERROR(RSQ($J${first+1}:$J${last},{col}{first+1}:{col}{last}),\"n/a\")", F_FX, "0.00")
    put(ws, f"A{r+1}", "Descriptive only: same-year macro data are published after the fact. Blank cells (e.g. unpublished years) "
                       "are skipped.", F_S)
    return ws, first, last


def build_inputs(wb, rates_last):
    ws = wb.create_sheet("Inputs")
    ws.column_dimensions["A"].width = 34
    for c, w in zip("BCDE", (14, 14, 14, 60)):
        ws.column_dimensions[c].width = w
    put(ws, "A1", "Inputs and assumptions", F_T)
    put(ws, "A2", "Blue = input. Shaded = override (blank uses latest data).", F_S)
    hdr(ws, 4, ["Item", "Latest data", "Override", "Used", "Source / note"])
    rows = [
        ("KRW rate (ACT/365)", f"=Rates_Monthly!B{rates_last}/100", "Latest OECD 3-month rate (monthly). Desks would use CD91 or KRW "
                                                                      "deposit/KOFR rates: type them in Override."),
        ("USD rate (ACT/360)", f"=Rates_Monthly!C{rates_last}/100", "Latest OECD 3-month rate (monthly). Override with SOFR term rate "
                                                                      "if preferred."),
    ]
    for i, (lab, f, note) in enumerate(rows):
        r = 5 + i
        put(ws, f"A{r}", lab, F_B)
        put(ws, f"B{r}", f, F_LINK, PCT2)
        put(ws, f"C{r}", None, F_IN, PCT2, fill=FILL_KEY)
        put(ws, f"D{r}", f'=IF(C{r}="",B{r},C{r})', F_FX, PCT2)
        put(ws, f"E{r}", note, F_S)
    put(ws, "A7", "Rate data as of", F_B)
    put(ws, "B7", f"=Rates_Monthly!A{rates_last}", F_LINK, "yyyy-mm")
    put(ws, "A9", "Hedge notional (USD)", F_B)
    put(ws, "D9", 10_000_000, F_IN, '#,##0', fill=FILL_KEY)
    put(ws, "E9", "Used for the hedge-cost example on the Dashboard.", F_S)
    put(ws, "A10", "Trading days per year", F_B)
    put(ws, "D10", 252, F_IN, '0')
    put(ws, "E10", "Annualization for realized volatility.", F_S)
    hdr(ws, 12, ["Forward tenor", "Days", "", "", "Note"])
    for i, (t, dd) in enumerate([("1M", 30), ("3M", 91), ("6M", 182), ("1Y", 365)]):
        r = 13 + i
        put(ws, f"A{r}", t, F_B)
        put(ws, f"B{r}", dd, F_IN, '0')
    put(ws, "E13", "Calendar days. Edit to match actual value dates.", F_S)
    hdr(ws, 18, ["Lookback", "Trading days", "", "", "Note"])
    for i, (t, dd) in enumerate([("1 day", 1), ("1 week", 5), ("1 month", 21), ("3 months", 63), ("1 year", 252)]):
        r = 19 + i
        put(ws, f"A{r}", t, F_B)
        put(ws, f"B{r}", dd, F_IN, '0')
    put(ws, "E19", "Observations back from the latest fixing.", F_S)
    return ws


def build_dashboard(wb, n_daily, extra_col, fm_first=5, fm_last=5, n_fd=0, d=None):
    ws = wb["Dashboard"]
    for c, w in zip("ABCDEFGHI", (30, 14, 14, 14, 14, 3, 30, 14, 14)):
        ws.column_dimensions[c].width = w
    L = "$B$3+1"  # row of the latest observation on Daily
    put(ws, "A1", "USD/KRW Monitor", F_T)
    put(ws, "A2", '="As of "&TEXT(B4,"d mmm yyyy")&"   |   Fed H.10, FRED, BOK ECOS, KRX"', F_S)
    for c in "ABCDEFGHI":
        ws[f"{c}2"].border = Border(bottom=MED)
    put(ws, "A3", "Observations (daily)"); put(ws, "B3", "=COUNT(Daily!$B:$B)", F_LINK, '#,##0')
    put(ws, "A4", "Last fixing"); put(ws, "B4", f"=INDEX(Daily!$A:$A,{L})", F_LINK, "d mmm yyyy")
    put(ws, "A5", "USD/KRW spot", F_B)
    put(ws, "B5", f"=INDEX(Daily!$B:$B,{L})", Font(name=FONT, size=11, bold=True), NUM2)

    section_at(ws, 7, 1, 5, "Spot performance (+ = weaker KRW)")
    hdr(ws, 8, ["Period", "Level then", "Change", "", ""])
    for i, lab in enumerate(["1 day", "1 week", "1 month", "3 months", "1 year"]):
        r = 9 + i
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", f"=INDEX(Daily!$B:$B,{L}-Inputs!$B${19+i})", F_LINK, NUM2)
        put(ws, f"C{r}", f"=$B$5/B{r}-1", F_FX, SGN)
    put(ws, "A14", "Year to date")
    put(ws, "B14", "=INDEX(Daily!$B:$B,MATCH(DATE(YEAR($B$4)-1,12,31),Daily!$A:$A,1))", F_LINK, NUM2)
    put(ws, "C14", "=$B$5/B14-1", F_FX, SGN)

    section_at(ws, 16, 1, 5, "Realized risk")
    hdr(ws, 17, ["Measure", "Value", "", "", ""])
    vol = lambda k: f"=STDEV(INDEX(Daily!$C:$C,{L}-{k-1}):INDEX(Daily!$C:$C,{L}))*SQRT(Inputs!$D$10)"
    rng = lambda fn: f"={fn}(INDEX(Daily!$B:$B,{L}-251):INDEX(Daily!$B:$B,{L}))"
    risk = [("Realized vol, 20 days (ann.)", vol(20), '0.0%'), ("Realized vol, 60 days (ann.)", vol(60), '0.0%'),
            ("Realized vol, 1 year (ann.)", vol(252), '0.0%'), ("1-year high", rng("MAX"), NUM2), ("1-year low", rng("MIN"), NUM2),
            ("Position in 1-year range", "=IFERROR(($B$5-B22)/(B21-B22),\"\")", '0%'),
            ("Z-score vs. 1-year mean", f"=($B$5-AVERAGE(INDEX(Daily!$B:$B,{L}-251):INDEX(Daily!$B:$B,{L})))"
                                       f"/STDEV(INDEX(Daily!$B:$B,{L}-251):INDEX(Daily!$B:$B,{L}))", '0.00'),
            ("Largest 1-day move, last year", f"=MAX(MAX(INDEX(Daily!$C:$C,{L}-251):INDEX(Daily!$C:$C,{L})),-MIN(INDEX(Daily!$C:$C,{L}-251):INDEX(Daily!$C:$C,{L})))", '0.00%')]
    for i, (lab, f, fmt) in enumerate(risk):
        r = 18 + i
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", f, F_LINK, fmt)

    section_at(ws, 27, 1, 5, "Implied forwards (CIP)")
    put(ws, "A28", "KRW rate used", F_B); put(ws, "B28", "=Inputs!D5", F_LINK, PCT2)
    put(ws, "C28", "USD rate used", F_B); put(ws, "D28", "=Inputs!D6", F_LINK, PCT2)
    hdr(ws, 29, ["Tenor", "Days", "Forward", "Swap points (won)", "Hedge carry (ann.)"])
    for i in range(4):
        r = 30 + i
        put(ws, f"A{r}", f"=Inputs!A{13+i}", F_LINK)
        put(ws, f"B{r}", f"=Inputs!B{13+i}", F_LINK, '0')
        put(ws, f"C{r}", f"=$B$5*(1+$B$28*B{r}/365)/(1+$D$28*B{r}/360)", F_FX, NUM2)
        put(ws, f"D{r}", f"=C{r}-$B$5", F_FX, NUM2)
        put(ws, f"E{r}", f"=(C{r}/$B$5-1)*365/B{r}", F_FX, PCT2)
    put(ws, "A34", "3M points x notional (KRW)", F_B)
    put(ws, "B34", "=Inputs!D9*D31", F_FX, '#,##0;(#,##0);-',
        comment="Selling USD forward against KRW: negative = the hedger receives fewer won than at spot.")
    put(ws, "A35", "Rate-implied; excludes the cross-currency basis.", F_S)

    section_at(ws, 7, 7, 3, "Asia FX and risk context")
    hdr(ws, 8, ["Measure", "Value", "1M change"], col=7)
    ctx = [("USD/JPY", "E", NUM2), ("USD/CNY", "G", NUM4)]
    r = 9
    for lab, col, fmt in ctx:
        put(ws, f"G{r}", lab)
        put(ws, f"H{r}", f'=IFERROR(LOOKUP(2,1/(Daily!${col}$2:${col}${n_daily+1}<>""),Daily!${col}$2:${col}${n_daily+1}),"n/a")', F_LINK, fmt)
        put(ws, f"I{r}", f'=IFERROR(H{r}/INDEX(Daily!${col}:${col},{L}-21)-1,"n/a")', F_FX, PCT2)
        r += 1
    for lab, key, fmt in [("USD broad index", "USD broad index", NUM2), ("VIX", "VIX", NUM2), ("UST 10Y %", "UST 10Y %", NUM2),
                          ("SOFR %", "SOFR %", NUM2)]:
        col = extra_col.get(key)
        put(ws, f"G{r}", lab)
        if col:
            put(ws, f"H{r}", f'=IFERROR(LOOKUP(2,1/(Daily!${col}$2:${col}${n_daily+1}<>""),Daily!${col}$2:${col}${n_daily+1}),"n/a")',
                F_LINK, fmt)
        else:
            put(ws, f"H{r}", "n/a", F_S, comment="Not loaded: run the program with internet access (FRED).")
            ws[f"H{r}"].alignment = Alignment(horizontal="right")
        r += 1
    section_at(ws, 16, 7, 3, "Correlation with KRW (daily returns)")
    hdr(ws, 17, ["Pair", "60 days", "1 year"], col=7)
    for i, (lab, col) in enumerate([("KRW vs JPY", "F"), ("KRW vs CNY", "H")]):
        rr = 18 + i
        put(ws, f"G{rr}", lab)
        for k, cc in [(60, "H"), (252, "I")]:
            put(ws, f"{cc}{rr}", f'=IFERROR(CORREL(INDEX(Daily!$C:$C,{L}-{k-1}):INDEX(Daily!$C:$C,{L}),'
                                 f'INDEX(Daily!${col}:${col},{L}-{k-1}):INDEX(Daily!${col}:${col},{L})),"n/a")', F_LINK, '0.00')
    put(ws, "G20", "Beta of KRW to JPY (1 year)")
    put(ws, "H20", f'=IFERROR(SLOPE(INDEX(Daily!$C:$C,{L}-251):INDEX(Daily!$C:$C,{L}),'
                   f'INDEX(Daily!$F:$F,{L}-251):INDEX(Daily!$F:$F,{L})),"n/a")', F_LINK, '0.00')
    put(ws, "G22", "Rate gap, KRW minus USD (pts)", F_B)
    put(ws, "H22", "=(B28-D28)*100", F_FX, NUM2)

    section_at(ws, 24, 7, 3, "Dollar supply and demand (US$ mn)")
    hdr(ws, 25, ["Measure", "Value", ""], col=7)
    fa = lambda c: f"Flows_Monthly!${c}${fm_first}:${c}${fm_last}"
    put(ws, "G26", "Latest balance-of-payments month")
    put(ws, "H26", f'=IFERROR(LOOKUP(2,1/({fa("B")}<>""),{fa("A")}),"n/a")', F_LINK, "mmm yyyy")
    for i, (lab, c) in enumerate([("Current account, that month", "B"), ("Basic balance, that month", "J"),
                                  ("Current account, 12M", "Q"), ("Residents' outward investment, 12M", "R"),
                                  ("Foreigners' inward investment, 12M", "S"), ("Basic balance, 12M", "T")]):
        r = 27 + i
        put(ws, f"G{r}", lab)
        put(ws, f"H{r}", f'=IFERROR(INDEX({fa(c)},MATCH($H$26,{fa("A")},0)),"n/a")', F_LINK, '#,##0;(#,##0);-')
    nf = "COUNT(Flows_Daily!$A:$A)+1"
    for i, k in enumerate((5, 20)):
        r = 33 + i
        put(ws, f"G{r}", f"Foreign equity net buying, {k}D (KRW bn)")
        put(ws, f"H{r}", f'=IF({nf}<{k + 1},"n/a",SUM(INDEX(Flows_Daily!$D:$D,{nf}-{k - 1}):INDEX(Flows_Daily!$D:$D,{nf})))',
            F_LINK, '#,##0;(#,##0);-')
    put(ws, "G35", "Corr. foreign buying vs KRW return, 1Y")
    fc = extra_col.get("Foreign net buy (KRW bn)")
    if fc:
        put(ws, "H35", f'=IFERROR(CORREL(INDEX(Daily!$C:$C,{L}-251):INDEX(Daily!$C:$C,{L}),'
                       f'INDEX(Daily!${fc}:${fc},{L}-251):INDEX(Daily!${fc}:${fc},{L})),"n/a")', F_LINK, '0.00')
    else:
        put(ws, "H35", "n/a", F_S)
    for r in range(26, 36):
        ws[f"H{r}"].alignment = Alignment(horizontal="right")
    put(ws, "G36", "Positive = net dollar supply.", F_S)

    # charts: last ~2 years of daily data, and monthly flows where available
    D = wb["Daily"]
    last_row = n_daily + 1
    first_row = max(2, last_row - 520)
    ch = new_chart("line", "USD/KRW, last 2 years")
    add_line(ch, D, 2, first_row, last_row, "USD/KRW", C_NAVY)
    ch.set_categories(Reference(D, min_col=1, min_row=first_row, max_row=last_row))
    _axes(ch, "#,##0", "mmm-yy", 126)
    ch.y_axis.scaling.min = None
    ws.add_chart(ch, "A39")
    ch2 = new_chart("line", "20-day realized volatility (annualized)")
    add_line(ch2, D, 4, first_row, last_row, "20D vol", C_RED)
    ch2.set_categories(Reference(D, min_col=1, min_row=first_row, max_row=last_row))
    _axes(ch2, "0%", "mmm-yy", 126)
    ws.add_chart(ch2, "G39")
    if n_fd and fm_last > fm_first:
        F = wb["Flows_Monthly"]
        r0 = max(fm_first, fm_last - 59)
        bar = new_chart("bar", "Foreign equity net buying (US$ mn, bars) vs USD/KRW (line, rhs)")
        add_bars(bar, F, 11, r0, fm_last, "Foreign equity net buying", C_GREY)
        bar.set_categories(Reference(F, min_col=1, min_row=r0, max_row=fm_last))
        _axes(bar, "#,##0", "mmm-yy", 12)
        ln = new_chart("line", "")
        add_line(ln, F, 13, r0, fm_last, "USD/KRW month-end", C_NAVY, 1.75)
        combo(bar, ln, "#,##0")
        if d is not None:
            k = d["usdkrw"][d["usdkrw"].index >= d["usdkrw"].index.max() - pd.DateOffset(months=61)]
            ln.y_axis.scaling.min = float(np.floor(k.min() / 50) * 50 - 50)
            ln.y_axis.scaling.max = float(np.ceil(k.max() / 50) * 50 + 50)
        legend_bottom(bar)
        ws.add_chart(bar, "A56")
    return ws


def build_stress(wb, n_daily):
    ws = wb.create_sheet("Stress")
    for c, w in zip("ABCDEFG", (6, 12, 12, 3, 6, 12, 12)):
        ws.column_dimensions[c].width = w
    put(ws, "A1", "Largest daily moves in USD/KRW", F_T)
    put(ws, "A2", "Log returns of the Fed H.10 daily rate. Ties return the first matching date.", F_S)
    rng = f"Daily!$C$3:$C${n_daily+1}"
    dts = f"Daily!$A$3:$A${n_daily+1}"
    hdr(ws, 4, ["#", "Date", "Won weaker"], col=1)
    hdr(ws, 4, ["#", "Date", "Won stronger"], col=5)
    for i in range(1, 16):
        r = 4 + i
        put(ws, f"A{r}", i, F_IN, '0')
        put(ws, f"C{r}", f"=LARGE({rng},A{r})", F_LINK, '0.00%')
        put(ws, f"B{r}", f"=INDEX({dts},MATCH(C{r},{rng},0))", F_LINK, "yyyy-mm-dd")
        put(ws, f"E{r}", i, F_IN, '0')
        put(ws, f"G{r}", f"=SMALL({rng},E{r})", F_LINK, '0.00%')
        put(ws, f"F{r}", f"=INDEX({dts},MATCH(G{r},{rng},0))", F_LINK, "yyyy-mm-dd")
    return ws


def build_flows_daily(wb, d):
    ws = wb.create_sheet("Flows_Daily")
    hdr(ws, 1, ["Date", "KOSPI foreign net (KRW bn)", "KOSDAQ foreign net (KRW bn)", "Total (KRW bn)", "Cumulative 20D"],
        widths=[11, 14, 14, 12, 13])
    ws.freeze_panes = "B2"
    krx = d.get("krx")
    if krx is None:
        put(ws, "A3", "No KRX data loaded. KRX now requires a free data.krx.co.kr account: pip install -U pykrx, set KRX_ID and KRX_PW, re-run (or pass --krx-csv).", F_S)
        return ws, 0
    for i, (dt, row) in enumerate(krx.iterrows()):
        r = i + 2
        put(ws, f"A{r}", pd.Timestamp(dt).to_pydatetime(), F_IN, "yyyy-mm-dd")
        for col, key in (("B", "kospi"), ("C", "kosdaq")):
            v = row.get(key)
            put(ws, f"{col}{r}", None if v is None or pd.isna(v) else float(v), F_IN, '#,##0.0;(#,##0.0);-')
        put(ws, f"D{r}", f"=SUM(B{r}:C{r})", F_FX, '#,##0.0;(#,##0.0);-')
        if r >= 21:
            put(ws, f"E{r}", f"=SUM(D{r-19}:D{r})", F_FX, '#,##0;(#,##0);-')
    return ws, len(krx)


def build_flows_monthly(wb, d, n_fd):
    ws = wb.create_sheet("Flows_Monthly")
    bop, krx = d.get("bop"), d.get("krx")
    labels = ["Month", "Current account", "Goods", "Direct inv. assets (out)", "Direct inv. liab. (in)", "Portfolio assets (out)",
              "Portfolio liab. (in)", "Residents' outward inv.", "Foreigners' inward inv.", "Basic balance",
              "Foreign equity net buy", "USD/KRW avg", "USD/KRW month-end", "KRW chg 1M (log)", "Rate gap (pts)",
              "Chg in rate gap", "Current acct 12M", "Outward inv. 12M", "Inward inv. 12M", "Basic balance 12M",
              "Foreign equity 12M", "KRW chg 12M (log)"]
    put(ws, "A1", "Korea FX flows (US$ mn)", F_T)
    put(ws, "A2", "Basic balance = current account + inward - outward investment (direct + portfolio, BPM6).", F_S)
    hdr(ws, 4, labels, widths=[10] + [12] * (len(labels) - 1))
    ws.row_dimensions[4].height = 42
    ws.freeze_panes = "B5"
    starts = [s.index.min() for s in (bop, krx) if s is not None]
    if not starts:
        put(ws, "A5", "No flow data loaded. Run with --ecos-key (balance of payments) and/or pykrx / --krx-csv (equity flows).", F_S)
        return ws, 5, 5
    m0 = max(pd.Timestamp(min(starts)).to_period("M").to_timestamp(), pd.Timestamp("1995-01-01"))
    m1 = pd.Timestamp(d["usdkrw"].index.max()).to_period("M").to_timestamp()
    months = pd.date_range(m0, m1, freq="MS")
    first = 5
    money = '#,##0;(#,##0);-'
    for i, mth in enumerate(months):
        r = first + i
        put(ws, f"A{r}", mth.to_pydatetime(), F_IN, "yyyy-mm")
        for col, key in zip("BCDEFG", ("ca", "goods", "di_a", "di_l", "pi_a", "pi_l")):
            v = bop[key].get(mth) if bop is not None and key in bop.columns else None
            put(ws, f"{col}{r}", None if v is None or pd.isna(v) else float(v), F_IN, money)
        put(ws, f"H{r}", f'=IF(COUNT(D{r},F{r})<2,"",D{r}+F{r})', F_FX, money)
        put(ws, f"I{r}", f'=IF(COUNT(E{r},G{r})<2,"",E{r}+G{r})', F_FX, money)
        put(ws, f"J{r}", f'=IF(COUNT(B{r},H{r},I{r})<3,"",B{r}+I{r}-H{r})', F_FX, money)
        if n_fd:
            fd = f"Flows_Daily!$A$2:$A${n_fd+1}"
            put(ws, f"K{r}", f'=IF(COUNTIFS({fd},">="&A{r},{fd},"<"&EDATE(A{r},1))=0,"",'
                             f'SUMIFS(Flows_Daily!$D$2:$D${n_fd+1},{fd},">="&A{r},{fd},"<"&EDATE(A{r},1))*1000/L{r})',
                F_LINK, money)
        put(ws, f"L{r}", f'=IFERROR(AVERAGEIFS(Daily!$B:$B,Daily!$A:$A,">="&A{r},Daily!$A:$A,"<"&EDATE(A{r},1)),"")', F_LINK, NUM1)
        put(ws, f"M{r}", f'=IFERROR(INDEX(Daily!$B:$B,MATCH(EDATE(A{r},1)-1,Daily!$A:$A,1)),"")', F_LINK, NUM1)
        put(ws, f"O{r}", f'=IFERROR(INDEX(Rates_Monthly!$D:$D,MATCH(A{r},Rates_Monthly!$A:$A,0)),"")', F_LINK, NUM2)
        if r > first:
            put(ws, f"N{r}", f'=IF(OR(M{r}="",M{r-1}=""),"",LN(M{r}/M{r-1}))', F_FX, PCT)
            put(ws, f"P{r}", f'=IF(OR(O{r}="",O{r-1}=""),"",O{r}-O{r-1})', F_FX, NUM2)
        if r >= first + 11:
            for src, dst in zip("BHIJK", "QRSTU"):
                put(ws, f"{dst}{r}", f'=IF(COUNT({src}{r-11}:{src}{r})<12,"",SUM({src}{r-11}:{src}{r}))', F_FX, money)
        if r >= first + 12:
            put(ws, f"V{r}", f'=IF(OR(M{r}="",M{r-12}=""),"",LN(M{r}/M{r-12}))', F_FX, PCT)
    last = first + len(months) - 1

    # ---- analysis block: which flow lines up with the won? (monthly, contemporaneous)
    c0 = 24  # column X
    X = lambda k: get_column_letter(c0 + k)
    section_at(ws, 4, c0, 5, "USD/KRW change vs. flows: one-variable regressions")
    hdr(ws, 5, ["Driver", "Months", "Correlation", "R²", "KRW % per $10bn"], col=c0, widths=[34, 9, 11, 9, 14])
    N, V = f"$N${first+1}:$N${last}", f"$V${first+12}:$V${last}"
    rows = [("1M: change in rate gap (per 1 pt)", "P", N, first + 1, 1),
            ("1M: current account", "B", N, first + 1, 10000), ("1M: residents' outward investment", "H", N, first + 1, 10000),
            ("1M: foreigners' inward investment", "I", N, first + 1, 10000), ("1M: basic balance", "J", N, first + 1, 10000),
            ("1M: foreign equity net buying", "K", N, first + 1, 10000),
            ("12M: current account", "Q", V, first + 12, 10000), ("12M: outward investment", "R", V, first + 12, 10000),
            ("12M: basic balance", "T", V, first + 12, 10000), ("12M: foreign equity net buying", "U", V, first + 12, 10000)]
    for i, (lab, col, y, r0, scale) in enumerate(rows):
        r = 6 + i
        x = f"${col}${r0}:${col}${last}"
        put(ws, f"{X(0)}{r}", lab)
        put(ws, f"{X(1)}{r}", f"=SUMPRODUCT(ISNUMBER({x})*ISNUMBER({y}))", F_FX, '0')
        put(ws, f"{X(2)}{r}", f'=IF({X(1)}{r}<12,"n/a",CORREL({y},{x}))', F_FX, '0.00')
        put(ws, f"{X(3)}{r}", f'=IF({X(1)}{r}<12,"n/a",RSQ({y},{x}))', F_FX, '0.00')
        put(ws, f"{X(4)}{r}", f'=IF({X(1)}{r}<12,"n/a",SLOPE({y},{x})*{scale})', F_FX, PCT2)
    put(ws, f"{X(4)}5", "KRW chg per $10bn")
    notes = ["KRW change > 0 = weaker won. A negative slope means more of that flow goes with a stronger won.",
             "Contemporaneous and descriptive: flows and prices are decided together, so this is association, not causation.",
             "12M rows use overlapping windows, which overstates R² and the effective number of observations.",
             "BoP figures are published about 5 weeks after month-end; foreign equity flows are same-day."]
    for i, t in enumerate(notes):
        put(ws, f"{X(0)}{18+i}", t, F_S)

    return ws, first, last


def build_charts(wb, d, n_daily, n_rates, n_fd, fm_first, fm_last):
    ws = wb.create_sheet("Charts")
    put(ws, "A1", "USD/KRW chart pack", F_T)
    put(ws, "A2", '=Dashboard!A2', F_S)
    for c in range(1, 20):
        ws.cell(2, c).border = Border(bottom=MED)
    D, R = wb["Daily"], wb["Rates_Monthly"]
    slots = iter([f"{c}{r}" for r in (4, 21, 38, 55, 72) for c in ("A", "K")])
    lr = n_daily + 1

    ch = new_chart("line", "USD/KRW since 1995 (monthly average)")
    add_line(ch, R, 5, 2, n_rates + 1, "USD/KRW", C_NAVY)
    ch.set_categories(Reference(R, min_col=1, min_row=2, max_row=n_rates + 1))
    _axes(ch, "#,##0", "yyyy", 60)
    ch.y_axis.scaling.min = float(np.floor(d["usdkrw"].min() / 100) * 100)
    ws.add_chart(ch, next(slots))

    f0 = max(2, lr - 520)
    ch = new_chart("line", "USD/KRW (lhs) vs USD/JPY (rhs), last 2 years")
    add_line(ch, D, 2, f0, lr, "USD/KRW", C_NAVY)
    ch.set_categories(Reference(D, min_col=1, min_row=f0, max_row=lr))
    _axes(ch, "#,##0", "mmm-yy", 126)
    j = new_chart("line", "")
    add_line(j, D, 5, f0, lr, "USD/JPY", C_RED)
    combo(ch, j, "0")
    legend_bottom(ch)
    ws.add_chart(ch, next(slots))

    ch = new_chart("line", "3-month rates: KRW vs USD (%)")
    add_line(ch, R, 2, 2, n_rates + 1, "KRW 3M", C_NAVY)
    add_line(ch, R, 3, 2, n_rates + 1, "USD 3M", C_RED)
    ch.set_categories(Reference(R, min_col=1, min_row=2, max_row=n_rates + 1))
    _axes(ch, "0.0", "yyyy", 60)
    legend_bottom(ch)
    ws.add_chart(ch, next(slots))

    ch = new_chart("line", "Rate-implied 3M hedge carry for a USD asset (ann.)")
    add_line(ch, R, 7, 2, n_rates + 1, "Hedge carry", C_BLACK)
    ch.set_categories(Reference(R, min_col=1, min_row=2, max_row=n_rates + 1))
    _axes(ch, "0.0%", "yyyy", 60)
    ws.add_chart(ch, next(slots))

    ch = new_chart("line", "20-day realized volatility since 1995")
    add_line(ch, D, 4, 22, lr, "20D vol", C_RED, 0.75)
    ch.set_categories(Reference(D, min_col=1, min_row=22, max_row=lr))
    _axes(ch, "0%", "yyyy", 756)
    ws.add_chart(ch, next(slots))

    if d.get("bop") is not None and fm_last > fm_first + 12:
        F = wb["Flows_Monthly"]
        r0 = fm_first + 11
        ch = new_chart("line", "12-month rolling flows (US$ mn)")
        for col, name, colour in ((17, "Current account", C_NAVY), (18, "Residents' outward inv.", C_RED),
                                  (19, "Foreigners' inward inv.", C_GREEN), (20, "Basic balance", C_BLACK)):
            add_line(ch, F, col, r0, fm_last, name, colour)
        ch.set_categories(Reference(F, min_col=1, min_row=r0, max_row=fm_last))
        _axes(ch, "#,##0", "yyyy", 60)
        legend_bottom(ch)
        ws.add_chart(ch, next(slots))

        bar = new_chart("bar", "12M basic balance (US$ mn, bars) vs 12M USD/KRW change (line, rhs)")
        add_bars(bar, F, 20, r0 + 1, fm_last, "Basic balance, 12M", C_GREY)
        bar.set_categories(Reference(F, min_col=1, min_row=r0 + 1, max_row=fm_last))
        _axes(bar, "#,##0", "yyyy", 60)
        ln = new_chart("line", "")
        add_line(ln, F, 22, r0 + 1, fm_last, "USD/KRW 12M change (+ = weaker KRW)", C_NAVY, 1.5)
        combo(bar, ln, "0%")
        legend_bottom(bar)
        ws.add_chart(bar, next(slots))

    if n_fd:
        FD = wb["Flows_Daily"]
        last = n_fd + 1
        r0 = max(21, last - 500)
        ch = new_chart("line", "Foreign equity net buying, 20-day cumulative (KRW bn)")
        add_line(ch, FD, 5, r0, last, "20D cumulative", C_NAVY)
        ch.set_categories(Reference(FD, min_col=1, min_row=r0, max_row=last))
        _axes(ch, "#,##0", "mmm-yy", 126)
        ws.add_chart(ch, next(slots))
    if n_fd and fm_last > fm_first:
        F = wb["Flows_Monthly"]
        r0 = max(fm_first + 11, fm_last - 119)
        ch = new_chart("bar", "Foreign equity net buying, monthly (US$ mn)")
        add_bars(ch, F, 11, r0, fm_last, "Monthly", C_NAVY)
        ch.set_categories(Reference(F, min_col=1, min_row=r0, max_row=fm_last))
        _axes(ch, "#,##0", "yyyy", 12)
        ws.add_chart(ch, next(slots))
    return ws


def build_method(wb, args):
    ws = wb.create_sheet("Methodology")
    for c, w in zip("ABCDEFG", (14, 46, 34, 9, 12, 12, 30)):
        ws.column_dimensions[c].width = w
    put(ws, "A1", "Methodology and sources", F_T)
    put(ws, "A2", f"Built {datetime.now():%Y-%m-%d %H:%M} by krw_fx_monitor.py"
                  f"{' (offline snapshot mode)' if args.offline else ''}.", F_S)
    hdr(ws, 4, ["Series", "Description", "Source used", "Obs", "First", "Last", "Note"])
    for i, (name, desc, src, n, f, l, note) in enumerate(LOG):
        r = 5 + i
        vals = [name, desc, src, n, f, l, note]
        for j, v in enumerate(vals):
            if isinstance(v, (pd.Timestamp, datetime)):
                v = pd.Timestamp(v).to_pydatetime()
            elif isinstance(v, np.integer):
                v = int(v)
            c = ws.cell(r, j + 1, v)
            c.font = Font(name=FONT, size=9, color=RED if src == "unavailable" else "000000")
            if j in (4, 5) and isinstance(v, datetime):
                c.number_format = "yyyy-mm-dd"
    r = 6 + len(LOG)
    notes = [
        "Colour code: blue = raw data / inputs; black = formulas; green = links to other sheets (shown black on the Dashboard); shaded = overrides.",
        "USD/KRW: Fed H.10 noon New York buying rate. Seoul closing rates (15:30 KST) differ slightly; add an ECOS key for BOK data.",
        "Implied forwards: covered interest parity, KRW ACT/365 and USD ACT/360. Market swap points also reflect the cross-currency "
        "basis and dollar funding conditions, so they can differ from rate-implied values.",
        "Monthly 3M rates are OECD interbank series via FRED; they lag and are proxies for desk funding rates. Override on Inputs.",
        "Realized volatility: standard deviation of daily log returns, annualized with the trading-day input.",
        "Macro R²: one-variable fits of Dec-to-Dec log change in USD/KRW on same-year macro gaps. Descriptive, not predictive.",
        "ECOS stat codes used (731Y001/0000001, 817Y002/010502000, 722Y001/0101000) should be checked against the ECOS catalogue; "
        "the program skips any series that fails to load.",
        "Flows: monthly balance of payments (BPM6, US$ mn) from Bank of Korea ECOS, items matched by name (see the Series log). "
        "Fix a mismatch with flows_config.json, e.g. {\"ca\": [\"301Y013\", \"000000\"]}; list items with --ecos-list 301Y013.",
        "Foreign equity flows: KRX daily net buying by foreign investors (KOSPI + KOSDAQ), via pykrx or your own CSV. Seoul dates; "
        "the Daily sheet aligns them to Fed H.10 dates (US holidays drop out), Flows_Daily keeps every Seoul trading day.",
        "Not public: onshore client order flow, exporter forward hedging, offshore NDF positioning, and real-time BOK intervention "
        "(disclosed quarterly with a lag).",
        "Not investment advice. Data are provided as published by each source; check critical figures before use.",
    ]
    for i, t in enumerate(notes):
        put(ws, f"A{r+i}", t, Font(name=FONT, size=9))
    return ws


def finish(wb):
    """Presentation pass: compact rows, no gridlines, black values on the Dashboard, print setup."""
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100
        ws.sheet_properties.outlinePr = None
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    db = wb["Dashboard"]
    for row in db.iter_rows(min_row=3, max_row=40, max_col=9):
        for c in row:
            if c.font and c.font.color is not None and c.font.color.rgb in ("00008000", "FF008000"):
                c.font = Font(name=FONT, size=c.font.size or 9, bold=c.font.bold, color="000000")
    db.sheet_view.zoomScale = 110


def build(d, args):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Dashboard"
    _, n_daily, extra_col = build_daily(wb, d)
    _, n_rates = build_rates(wb, d)
    build_inputs(wb, n_rates + 1)
    build_macro(wb, d)
    _, n_fd = build_flows_daily(wb, d)
    _, fm_first, fm_last = build_flows_monthly(wb, d, n_fd)
    build_stress(wb, n_daily)
    build_method(wb, args)
    build_dashboard(wb, n_daily, extra_col, fm_first, fm_last, n_fd, d)
    build_charts(wb, d, n_daily, n_rates, n_fd, fm_first, fm_last)
    wb.move_sheet("Inputs", offset=1 - wb.sheetnames.index("Inputs"))
    wb.move_sheet("Charts", offset=1 - wb.sheetnames.index("Charts"))
    finish(wb)
    wb.save(args.out)
    return args.out


def main():
    p = argparse.ArgumentParser(description="Build a USD/KRW desk monitor workbook.")
    p.add_argument("--out", default=f"KRW_FX_Monitor_{date.today():%Y%m%d}.xlsx")
    p.add_argument("--start", default="1995-01-01", help="first date for daily series (YYYY-MM-DD)")
    p.add_argument("--offline", action="store_true", help="use only the bundled data snapshot")
    p.add_argument("--ecos-key", default=os.environ.get("ECOS_API_KEY"), help="Bank of Korea ECOS API key (optional)")
    p.add_argument("--krx-csv", help="CSV of daily foreign net buying, KRW bn: columns date,kospi,kosdaq")
    p.add_argument("--flows-start", default="2005-01-01", help="first date for KRX foreign flows")
    p.add_argument("--ecos-list", metavar="STAT_CODE", help="print the item list of an ECOS table and exit")
    args = p.parse_args()
    if args.ecos_list:
        if not args.ecos_key:
            sys.exit("--ecos-list needs --ecos-key")
        for code, name, cyc in ecos_items(args.ecos_key, args.ecos_list):
            print(f"{code:<12} {cyc:<6} {name}")
        return
    d = collect(args)
    if d["usdkrw"] is None or d["kr3m"] is None or d["us3m"] is None:
        sys.exit("USD/KRW or rate data unavailable from every source; cannot build the workbook.")
    out = build(d, args)
    print(f"Saved {out}")


if __name__ == "__main__":
    main()
